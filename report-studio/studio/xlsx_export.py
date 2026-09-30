"""Excel companion file: the report outline plus every data table, styled."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .docx_builder import LABELS
from .themes import get_theme


def _iter_tables(report: dict):
    for s in report.get("sections") or []:
        for holder in [s, *(s.get("subsections") or [])]:
            for b in holder.get("blocks") or []:
                if b.get("type") == "table":
                    yield holder["heading"], b


def build_xlsx(report: dict, out_path: str | Path) -> Path:
    rtl = report.get("lang") != "en"
    L = LABELS["ar" if rtl else "en"]
    th = get_theme(report.get("theme"))
    head_font = Font(name="Tajawal", bold=True, color="FFFFFF", size=12)
    body_font = Font(name="Amiri", size=12)
    fill = PatternFill("solid", fgColor=th.primary)
    zebra = PatternFill("solid", fgColor=th.soft)
    thin = Side(style="thin", color="C9CED6")
    box = Border(top=thin, bottom=thin, left=thin, right=thin)
    wrap = Alignment(wrap_text=True, vertical="center", horizontal="center")

    wb = Workbook()
    ws = wb.active
    ws.title = "المحتويات" if rtl else "Outline"
    ws.sheet_view.rightToLeft = rtl
    cover = report.get("cover") or {}
    ws["A1"] = cover.get("title", "")
    ws["A1"].font = Font(name="Tajawal", bold=True, size=16, color=th.primary)
    ws["A2"] = " | ".join(x for x in (cover.get("student"), cover.get("subject"),
                                     cover.get("university")) if x)
    ws["A2"].font = Font(name="Tajawal", size=11, color=th.muted)
    row = 4
    for i, s in enumerate(report.get("sections") or [], 1):
        ws.cell(row, 1, i).font = Font(name="Tajawal", bold=True, color=th.accent)
        ws.cell(row, 2, s["heading"]).font = Font(name="Tajawal", bold=True, size=12)
        row += 1
        for j, sub in enumerate(s.get("subsections") or [], 1):
            ws.cell(row, 1, f"{i}.{j}").font = Font(name="Tajawal", color=th.muted)
            ws.cell(row, 2, sub["heading"]).font = Font(name="Tajawal", size=11)
            row += 1
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 70

    for n, (section, t) in enumerate(_iter_tables(report), 1):
        sh = wb.create_sheet(f"{L['table']} {n}")
        sh.sheet_view.rightToLeft = rtl
        headers, rows = t.get("headers") or [], t.get("rows") or []
        ncols = max(len(headers), *(len(r) for r in rows)) if rows else len(headers)
        sh.cell(1, 1, t.get("caption") or section).font = Font(
            name="Tajawal", bold=True, size=14, color=th.primary)
        sh.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, ncols))
        for j, h in enumerate(headers, 1):
            c = sh.cell(3, j, h)
            c.font, c.fill, c.alignment, c.border = head_font, fill, wrap, box
        for i, r in enumerate(rows, 4):
            for j in range(1, ncols + 1):
                v = r[j - 1] if j - 1 < len(r) else ""
                c = sh.cell(i, j, _number(v))
                c.font, c.alignment, c.border = body_font, wrap, box
                if i % 2 == 1:
                    c.fill = zebra
        for j in range(1, ncols + 1):
            sh.column_dimensions[sh.cell(3, j).column_letter].width = 24
        sh.freeze_panes = "A4"

    out_path = Path(out_path)
    wb.save(out_path)
    return out_path


def _number(v: str):
    """Keep numbers numeric in Excel so students can chart them."""
    s = str(v).strip().replace(",", "")
    try:
        return int(s) if s.lstrip("-").isdigit() else float(s)
    except ValueError:
        return v
