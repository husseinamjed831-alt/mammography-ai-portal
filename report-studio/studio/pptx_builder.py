"""Build a designed presentation (.pptx) from structured slide content.

Three tiers:
  standard  clean themed slides
  animated  + fade transitions and staggered entrance animations
  three_d   + native PowerPoint 3-D shapes (extrusion, bevel, lighting) that
            glide and turn between slides with the Morph transition

Everything is plain DrawingML/PresentationML that PowerPoint 2019/365 opens
natively; Morph falls back to a fade on older versions. Layouts are designed
left-to-right and mirrored for Arabic.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .docx_builder import has_arabic
from .themes import Theme, get_theme

SW, SH = 13.333, 7.5  # 16:9 in inches
M = 0.7               # side margin
FONT = "Segoe UI"     # ships with Windows/Office and covers Arabic

NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
NS_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
NS_P159 = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"

TXT = {
    "ar": {"agenda": "محاور العرض", "thanks": "شكراً لإصغائكم", "questions": "هل من أسئلة؟",
           "by": "إعداد", "sup": "إشراف", "subject": "المادة"},
    "en": {"agenda": "Outline", "thanks": "Thank you", "questions": "Questions?",
           "by": "Presented by", "sup": "Supervised by", "subject": "Course"},
}


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def _mix(hex_a: str, hex_b: str, t: float) -> str:
    a = [int(hex_a[i:i + 2], 16) for i in (0, 2, 4)]
    b = [int(hex_b[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(a, b))


class DeckBuilder:
    def __init__(self, deck: dict):
        self.d = deck
        self.lang = "en" if deck.get("lang") == "en" else "ar"
        self.rtl = self.lang == "ar"
        self.T = TXT[self.lang]
        self.th: Theme = get_theme(deck.get("theme"))
        self.tier = deck.get("tier") or "standard"
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(SW), Inches(SH)
        self.blank = self.prs.slide_layouts[6]
        self.n = 0

    # ------------------------------------------------------------ geometry

    def X(self, x: float, w: float) -> float:
        """Mirror an LTR x-position for right-to-left decks."""
        return SW - x - w if self.rtl else x

    # ------------------------------------------------------------ primitives

    def box(self, slide, x, y, w, h, *, fill=None, shape=MSO_SHAPE.RECTANGLE, name=None,
            line=None, radius=None, shadow=False, mirror=True):
        sx = self.X(x, w) if mirror else x
        shp = slide.shapes.add_shape(shape, Inches(sx), Inches(y), Inches(w), Inches(h))
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(1)
        else:
            shp.line.fill.background()
        if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            shp.adjustments[0] = radius
        if name:
            shp.name = name
        if shadow:
            self._shadow(shp)
        else:
            # An explicit empty list stops the theme's default effects applying.
            etree.SubElement(shp._element.spPr, qn("a:effectLst"))
        shp.text_frame.text = ""
        return shp

    def _shadow(self, shp):
        spPr = shp._element.spPr
        eff = etree.SubElement(spPr, qn("a:effectLst"))
        sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad="190500", dist="38100",
                              dir="5400000", algn="t", rotWithShape="0")
        clr = etree.SubElement(sh, qn("a:srgbClr"), val="1F2328")
        etree.SubElement(clr, qn("a:alpha"), val="16000")

    def text(self, slide, x, y, w, h, paras, *, size=18, color=None, bold=False,
             align="start", anchor="top", line=1.1, name=None, space_after=6,
             mirror=True, autofit=True):
        """paras: list of str or dicts {text, size, color, bold, bullet}."""
        sx = self.X(x, w) if mirror else x
        tb = slide.shapes.add_textbox(Inches(sx), Inches(y), Inches(w), Inches(h))
        if name:
            tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = Inches(0.03)
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                              "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        # Shrink-on-overflow keeps long AI-written text inside its box.
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE if autofit else MSO_AUTO_SIZE.NONE
        for i, p in enumerate(paras):
            spec = {"text": p} if isinstance(p, str) else p
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            a = spec.get("align", align)
            para.alignment = {
                "start": PP_ALIGN.RIGHT if self.rtl else PP_ALIGN.LEFT,
                "end": PP_ALIGN.LEFT if self.rtl else PP_ALIGN.RIGHT,
                "center": PP_ALIGN.CENTER,
                "justify": PP_ALIGN.JUSTIFY,
            }[a]
            para.line_spacing = spec.get("line", line)
            para.space_after = Pt(spec.get("space_after", space_after))
            if self.rtl and has_arabic(str(spec["text"])):
                para._p.get_or_add_pPr().set("rtl", "1")
            psize = spec.get("size", size)
            pcolor = spec.get("color", color or self.th.ink)
            if spec.get("bullet"):
                self._run(para, spec["bullet"] + "  ", psize * 0.8, self.th.accent, True)
            for j, chunk in enumerate(str(spec["text"]).split("**")):
                if chunk:
                    self._run(para, chunk, psize, pcolor, spec.get("bold", bold) or j % 2 == 1)
        return tb

    def _run(self, para, text, size, color, bold):
        r = para.add_run()
        r.text = text
        f = r.font
        f.size = Pt(size)
        f.bold = bold
        f.name = FONT
        f.color.rgb = rgb(color)
        rPr = r._r.get_or_add_rPr()
        rPr.set("lang", "ar-IQ" if self.rtl else "en-US")
        etree.SubElement(rPr, qn("a:cs"), typeface=FONT)
        return r

    # ------------------------------------------------------------ 3-D & decor

    def _make_3d(self, shp, *, preset="isometricOffAxis2Left", rot=None, depth=60,
                 bevel=38, material="softEdge"):
        spPr = shp._element.spPr
        scene = etree.SubElement(spPr, qn("a:scene3d"))
        cam = etree.SubElement(scene, qn("a:camera"), prst=preset)
        if rot:
            etree.SubElement(cam, qn("a:rot"), lat=str(rot[0] * 60000),
                             lon=str(rot[1] * 60000), rev=str(rot[2] * 60000))
        etree.SubElement(scene, qn("a:lightRig"), rig="threePt", dir="t")
        sp3d = etree.SubElement(spPr, qn("a:sp3d"), extrusionH=str(int(Pt(depth))),
                                prstMaterial=material)
        etree.SubElement(sp3d, qn("a:bevelT"), w=str(int(Pt(bevel))), h=str(int(Pt(bevel))))
        ext = etree.SubElement(sp3d, qn("a:extrusionClr"))
        etree.SubElement(ext, qn("a:srgbClr"), val=_mix(self.th.primary, "000000", 0.25))

    # Positions (LTR inches) the decorative objects take on successive slides.
    # Morph animates between them because every slide reuses the same "!!"
    # shape names. They stay in the corners, clear of titles and content.
    ORB_PATH = [(10.4, 4.6, 3.6), (11.0, -1.7, 3.0), (-1.7, 5.3, 3.0), (10.9, 5.0, 2.8),
                (11.4, -1.2, 2.4), (-1.4, -1.5, 3.0)]
    SLAB_PATH = [(11.9, 0.35, 0.9, 20), (12.3, 6.3, 1.1, -25), (11.7, 0.45, 0.75, 40),
                 (12.2, 6.1, 1.2, -10), (11.95, 0.3, 1.0, 55)]
    CUBE_PATH = [(12.6, 1.5, 0.4), (11.8, 6.75, 0.45), (12.55, 6.35, 0.35), (11.6, 1.45, 0.4),
                 (12.6, 5.85, 0.4)]

    def decor(self, slide, dark=False):
        k = self.n
        ox, oy, od = self.ORB_PATH[k % len(self.ORB_PATH)]
        orb_col = _mix(self.th.accent, "FFFFFF", 0.15 if dark else 0.82)
        orb = self.box(slide, ox, oy, od, od, fill=orb_col, shape=MSO_SHAPE.OVAL, name="!!orb")
        if self.tier != "three_d":
            return
        self._make_3d(orb, preset="orthographicFront", depth=0, bevel=od * 36, material="powder")
        sx, sy, ss, ang = self.SLAB_PATH[k % len(self.SLAB_PATH)]
        slab_col = _mix(self.th.primary, "FFFFFF", 0.35 if dark else 0.1)
        slab = self.box(slide, sx, sy, ss, ss, fill=slab_col, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                        name="!!slab", radius=0.18)
        slab.rotation = ang
        self._make_3d(slab, preset="isometricOffAxis2Left", depth=34, bevel=6, material="plastic")
        cx, cy, cs = self.CUBE_PATH[k % len(self.CUBE_PATH)]
        cube = self.box(slide, cx, cy, cs, cs, fill=self.th.accent, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                        name="!!cube", radius=0.2)
        cube.rotation = -ang * 1.5
        self._make_3d(cube, preset="isometricOffAxis1Right", depth=26, bevel=5, material="plastic")

    # ------------------------------------------------------------ chrome

    def new_slide(self, *, bg=None):
        slide = self.prs.slides.add_slide(self.blank)
        if bg:
            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = rgb(bg)
        self.n += 1
        return slide

    def header(self, slide, title):
        size = 30 if len(title) < 45 else 26 if len(title) < 70 else 22
        t = self.text(slide, M, 0.35, SW - 2 * M - 1.4, 0.95, [title], size=size, bold=True,
                      color=self.th.primary, anchor="bottom", name="Title")
        self.box(slide, M, 1.45, 1.1, 0.07, fill=self.th.accent, name="!!rule")
        return t

    def footer(self, slide, dark=False):
        col = "FFFFFF" if dark else self.th.muted
        self.text(slide, M, SH - 0.52, 8, 0.35, [self.d.get("cover", {}).get("title", "")],
                  size=10, color=col, autofit=False)
        self.text(slide, SW - M - 1, SH - 0.52, 1, 0.35, [str(self.n)], size=11, bold=True,
                  color=self.th.accent, align="end", autofit=False)

    # ------------------------------------------------------------ layouts

    def title_slide(self):
        c, th = self.d.get("cover") or {}, self.th
        s = self.new_slide(bg=th.primary)
        self.decor(s, dark=True)
        top = " — ".join(x for x in (c.get("university"), c.get("college")) if x)
        dept = c.get("department")
        self.text(s, M, 0.55, SW - 2 * M, 0.9, [top] + ([dept] if dept else []), size=15,
                  color=_mix(th.primary, "FFFFFF", 0.78), bold=True, space_after=2)
        title = c.get("title", "")
        size = 44 if len(title) < 40 else 38 if len(title) < 65 else 32
        self.text(s, M, 1.7, SW - 2 * M - 2.5, 2.2, [title], size=size, bold=True,
                  color="FFFFFF", anchor="bottom", line=1.05, name="Title")
        self.box(s, M, 4.08, 1.4, 0.08, fill=th.accent, name="!!rule")
        if c.get("subtitle"):
            self.text(s, M, 4.3, SW - 2 * M - 2.5, 0.7, [c["subtitle"]], size=18,
                      color=_mix(th.primary, "FFFFFF", 0.75), name="Sub")
        info = [(self.T["by"], c.get("student")), (self.T["sup"], c.get("supervisor")),
                (self.T["subject"], c.get("subject"))]
        info = [(k, v) for k, v in info if v]
        colw = (SW - 2 * M) / max(1, len(info))
        for i, (k, v) in enumerate(info):
            self.text(s, M + i * colw, 5.55, colw - 0.2, 1.1,
                      [{"text": k, "size": 12, "color": _mix(th.accent, "FFFFFF", 0.35), "bold": True},
                       {"text": v, "size": 17, "color": "FFFFFF", "bold": True}], space_after=2)
        if c.get("academic_year"):
            self.text(s, M, SH - 0.55, 3, 0.35, [c["academic_year"]], size=11,
                      color=_mix(th.primary, "FFFFFF", 0.6), autofit=False)
        return s, ["Title", "Sub"]

    def agenda_slide(self, titles):
        s = self.new_slide()
        self.decor(s)
        self.header(s, self.T["agenda"])
        rows = (len(titles) + 1) // 2
        rh = min(1.05, 4.9 / max(rows, 1))
        names = []
        for i, t in enumerate(titles):
            col, row = divmod(i, rows)
            x, y = M + col * 6.1, 1.9 + row * rh
            self.box(s, x, y + 0.08, 0.62, 0.62, fill=self.th.primary, shape=MSO_SHAPE.OVAL,
                     name=f"AgendaDot{i}")
            self.text(s, x, y + 0.08, 0.62, 0.62, [f"{i + 1:02d}"], size=15, bold=True,
                      color="FFFFFF", align="center", anchor="middle", autofit=False)
            nm = f"Agenda{i}"
            self.text(s, x + 0.85, y, 5.0, 0.78, [t], size=18, bold=True, anchor="middle", name=nm)
            names.append(nm)
        self.footer(s)
        return s, ["Title", *names]

    def section_slide(self, sl, idx):
        th = self.th
        s = self.new_slide(bg=th.primary)
        self.decor(s, dark=True)
        self.text(s, M, 1.7, 4, 1.6, [f"{idx:02d}"], size=96, bold=True,
                  color=_mix(th.accent, "FFFFFF", 0.2), name="Num")
        self.box(s, M, 3.55, 1.4, 0.08, fill=th.accent, name="!!rule")
        self.text(s, M, 3.8, SW - 2 * M - 2.5, 1.5, [sl["title"]], size=40, bold=True,
                  color="FFFFFF", name="Title")
        if sl.get("subtitle"):
            self.text(s, M, 5.2, SW - 2 * M - 2.5, 0.9, [sl["subtitle"]], size=18,
                      color=_mix(th.primary, "FFFFFF", 0.75), name="Sub")
        return s, ["Num", "Title", "Sub"] if sl.get("subtitle") else ["Num", "Title"]

    def bullets_slide(self, sl):
        s = self.new_slide()
        self.decor(s)
        self.header(s, sl["title"])
        items = sl.get("bullets") or []
        chars = sum(len(b) for b in items)
        size = 22 if chars < 260 else 19 if chars < 420 else 16
        hl = sl.get("highlight") or {}
        width = SW - 2 * M - (4.3 if hl.get("value") else 0.4)
        mark = "◂" if self.rtl else "▸"
        self.text(s, M, 1.85, width, 4.9, [{"text": b, "bullet": mark} for b in items],
                  size=size, space_after=14, line=1.15, name="Body")
        names = ["Title", "Body"]
        if hl.get("value"):
            x = SW - M - 3.8
            self.box(s, x, 1.95, 3.8, 3.9, fill=self.th.soft, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                     radius=0.08, shadow=True, name="Card")
            self.text(s, x + 0.3, 2.3, 3.2, 1.5, [hl["value"]], size=54, bold=True,
                      color=self.th.primary, align="center", anchor="middle", name="CardVal")
            self.text(s, x + 0.3, 3.85, 3.2, 1.8, [hl.get("label", "")], size=16,
                      color=self.th.ink, align="center", name="CardLbl")
            names += ["Card", "CardVal", "CardLbl"]
        self.footer(s)
        return s, names

    def two_column_slide(self, sl):
        s = self.new_slide()
        self.decor(s)
        self.header(s, sl["title"])
        names = ["Title"]
        cw = (SW - 2 * M - 0.5) / 2
        mark = "◂" if self.rtl else "▸"
        for i, key in enumerate(("left", "right")):
            col = sl.get(key) or {}
            x = M + i * (cw + 0.5)
            fill = self.th.soft if i == 0 else _mix(self.th.accent, "FFFFFF", 0.88)
            self.box(s, x, 1.85, cw, 4.9, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                     radius=0.06, shadow=True, name=f"Col{i}")
            self.box(s, x, 1.85, cw, 0.09, fill=self.th.primary if i == 0 else self.th.accent)
            items = col.get("bullets") or []
            size = 18 if sum(len(b) for b in items) < 260 else 15
            self.text(s, x + 0.35, 2.15, cw - 0.7, 4.4,
                      [{"text": col.get("heading", ""), "size": 22, "bold": True,
                        "color": self.th.primary, "space_after": 12}]
                      + [{"text": b, "bullet": mark} for b in items],
                      size=size, space_after=10, name=f"ColText{i}")
            names += [f"Col{i}", f"ColText{i}"]
        self.footer(s)
        return s, names

    def stats_slide(self, sl):
        s = self.new_slide()
        self.decor(s)
        self.header(s, sl["title"])
        stats = (sl.get("stats") or [])[:4]
        n = max(1, len(stats))
        gap = 0.35
        cw = (SW - 2 * M - gap * (n - 1)) / n
        names = ["Title"]
        for i, st in enumerate(stats):
            x = M + i * (cw + gap)
            self.box(s, x, 2.1, cw, 3.6, fill="FFFFFF", line=_mix(self.th.accent, "FFFFFF", 0.6),
                     shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07, shadow=True, name=f"Stat{i}")
            self.box(s, x, 2.1, cw, 0.1, fill=self.th.accent)
            val = str(st.get("value", ""))
            vsize = 50 if len(val) <= 5 else 40 if len(val) <= 8 else 30
            self.text(s, x + 0.2, 2.45, cw - 0.4, 1.4, [val], size=vsize, bold=True,
                      color=self.th.primary, align="center", anchor="middle", name=f"StatVal{i}")
            self.text(s, x + 0.25, 3.9, cw - 0.5, 1.7, [st.get("label", "")], size=16,
                      align="center", name=f"StatLbl{i}")
            names += [f"Stat{i}", f"StatVal{i}", f"StatLbl{i}"]
        if sl.get("note"):
            self.text(s, M, 6.0, SW - 2 * M, 0.6, [sl["note"]], size=13, color=self.th.muted,
                      align="center", name="Note")
            names.append("Note")
        self.footer(s)
        return s, names

    def table_slide(self, sl):
        s = self.new_slide()
        self.decor(s)
        self.header(s, sl["title"])
        t = sl.get("table") or {}
        headers, rows = t.get("headers") or [], t.get("rows") or []
        ncols = max(len(headers), *(len(r) for r in rows)) if rows else len(headers)
        nrows = 1 + len(rows)
        h = min(0.62 * nrows, 4.8)
        gf = s.shapes.add_table(nrows, ncols, Inches(M), Inches(1.95), Inches(SW - 2 * M), Inches(h))
        gf.name = "Table"
        tbl = gf.table
        tblPr = tbl._tbl.tblPr
        if self.rtl:
            tblPr.set("rtl", "1")
        for att in ("bandRow", "firstRow"):
            tblPr.attrib.pop(att, None)
        style = tblPr.find(qn("a:tableStyleId"))
        if style is not None:
            tblPr.remove(style)
        size = 15 if nrows <= 6 else 13
        for i in range(nrows):
            for j in range(ncols):
                cell = tbl.cell(i, j)
                src = headers if i == 0 else rows[i - 1]
                val = src[j] if j < len(src) else ""
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(self.th.primary if i == 0 else
                                               self.th.soft if i % 2 == 0 else "FFFFFF")
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.paragraphs[0].text = ""
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                if self.rtl and has_arabic(val):
                    p._p.get_or_add_pPr().set("rtl", "1")
                self._run(p, val, size, "FFFFFF" if i == 0 else self.th.ink, i == 0)
        if t.get("caption"):
            self.text(s, M, 1.95 + h + 0.15, SW - 2 * M, 0.5, [t["caption"]], size=12,
                      color=self.th.muted, align="center", name="Caption")
        self.footer(s)
        return s, ["Title", "Table"]

    def quote_slide(self, sl):
        th = self.th
        s = self.new_slide(bg=th.soft)
        self.decor(s)
        q = sl.get("quote") or {}
        self.text(s, 1.5, 0.9, SW - 3, 1.4, ["❝"], size=80, color=th.accent, align="center",
                  autofit=False, name="Mark")
        size = 32 if len(q.get("text", "")) < 120 else 26
        self.text(s, 1.5, 2.3, SW - 3, 2.8, [q.get("text", "")], size=size, bold=True,
                  color=th.primary, align="center", anchor="middle", line=1.2, name="Quote")
        self.text(s, 1.5, 5.3, SW - 3, 0.6, [f"— {q.get('source', '')}"], size=16,
                  color=th.muted, align="center", name="Source")
        self.footer(s)
        return s, ["Mark", "Quote", "Source"]

    def timeline_slide(self, sl):
        s = self.new_slide()
        self.decor(s)
        self.header(s, sl["title"])
        steps = (sl.get("steps") or [])[:6]
        n = max(1, len(steps))
        span = SW - 2 * M
        cw = span / n
        self.box(s, M + cw / 2, 2.62, span - cw, 0.05, fill=_mix(self.th.accent, "FFFFFF", 0.4))
        names = ["Title"]
        for i, st in enumerate(steps):
            cx = M + i * cw
            self.box(s, cx + cw / 2 - 0.36, 2.28, 0.72, 0.72, fill=self.th.primary,
                     shape=MSO_SHAPE.OVAL, shadow=True, name=f"Dot{i}")
            self.text(s, cx + cw / 2 - 0.36, 2.28, 0.72, 0.72, [str(i + 1)], size=18, bold=True,
                      color="FFFFFF", align="center", anchor="middle", autofit=False, name=f"DotN{i}")
            size = 14 if n >= 5 else 16
            self.text(s, cx + 0.12, 3.25, cw - 0.24, 3.3,
                      [{"text": st.get("title", ""), "size": size + 3, "bold": True,
                        "color": self.th.primary, "space_after": 6},
                       {"text": st.get("text", ""), "size": size}],
                      align="center", name=f"Step{i}")
            names += [f"Dot{i}", f"DotN{i}", f"Step{i}"]
        self.footer(s)
        return s, names

    def closing_slide(self):
        th, c = self.th, self.d.get("cover") or {}
        s = self.new_slide(bg=th.primary)
        self.decor(s, dark=True)
        self.text(s, M, 2.2, SW - 2 * M, 1.5, [self.T["thanks"]], size=54, bold=True,
                  color="FFFFFF", align="center", anchor="middle", name="Title")
        self.box(s, SW / 2 - 0.7, 3.85, 1.4, 0.08, fill=th.accent, name="!!rule", mirror=False)
        self.text(s, M, 4.15, SW - 2 * M, 0.8, [self.T["questions"]], size=24,
                  color=_mix(th.primary, "FFFFFF", 0.75), align="center", name="Sub")
        if c.get("student"):
            self.text(s, M, 5.4, SW - 2 * M, 0.6, [c["student"]], size=16, bold=True,
                      color=_mix(th.accent, "FFFFFF", 0.3), align="center", name="Name")
        return s, ["Title", "Sub"]

    # ------------------------------------------------------------ motion

    def transition(self, slide):
        if self.tier == "standard":
            return
        sld = slide._element
        if self.tier == "three_d":
            xml = (
                f'<mc:AlternateContent xmlns:mc="{NS_MC}">'
                f'<mc:Choice xmlns:p159="{NS_P159}" Requires="p159">'
                f'<p:transition xmlns:p="{NS_P}" xmlns:p14="{NS_P14}" spd="slow" p14:dur="1600">'
                '<p159:morph option="byObject"/></p:transition></mc:Choice>'
                f'<mc:Fallback><p:transition xmlns:p="{NS_P}" spd="slow"><p:fade/></p:transition>'
                "</mc:Fallback></mc:AlternateContent>"
            )
        else:
            xml = f'<p:transition xmlns:p="{NS_P}" spd="med"><p:fade/></p:transition>'
        el = etree.fromstring(xml)
        anchor = sld.find(qn("p:clrMapOvr"))
        (anchor if anchor is not None else sld.find(qn("p:cSld"))).addnext(el)

    def animate(self, slide, names: list[str]):
        """Staggered fade-in of the named shapes as soon as the slide opens."""
        if self.tier == "standard":
            return
        by_name = {sh.name: sh for sh in slide.shapes}
        targets = [by_name[n] for n in names if n in by_name]
        if not targets:
            return
        ids = iter(range(5, 10_000))  # 1-4 are the root, sequence and group nodes
        effects = []
        builds = []
        for k, shp in enumerate(targets):
            spid = shp.shape_id
            para_mode = shp.has_text_frame and shp.name == "Body" and len(shp.text_frame.paragraphs) > 1
            chunks = range(len(shp.text_frame.paragraphs)) if para_mode else [None]
            for pi in chunks:
                tgt = (f'<p:spTgt spid="{spid}"><p:txEl><p:pRg st="{pi}" end="{pi}"/></p:txEl></p:spTgt>'
                       if pi is not None else f'<p:spTgt spid="{spid}"/>')
                delay = 250 * len(effects)
                c1, c2, c3 = next(ids), next(ids), next(ids)
                effects.append(
                    f'<p:par><p:cTn id="{c1}" presetID="10" presetClass="entr" presetSubtype="0" '
                    f'fill="hold" grpId="0" nodeType="withEffect"><p:stCondLst><p:cond delay="{delay}"/>'
                    f'</p:stCondLst><p:childTnLst>'
                    f'<p:set><p:cBhvr><p:cTn id="{c2}" dur="1" fill="hold"><p:stCondLst>'
                    f'<p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl>{tgt}</p:tgtEl>'
                    f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>'
                    f'</p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
                    f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{c3}" dur="600"/>'
                    f'<p:tgtEl>{tgt}</p:tgtEl></p:cBhvr></p:animEffect>'
                    f'</p:childTnLst></p:cTn></p:par>'
                )
            if shp.has_text_frame and shp.text_frame.text.strip():
                build = ' build="p"' if para_mode else ""
                builds.append(f'<p:bldP spid="{spid}" grpId="0"{build}/>')
        xml = (
            f'<p:timing xmlns:p="{NS_P}"><p:tnLst><p:par>'
            '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
            '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
            '<p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst>'
            '<p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond>'
            '</p:stCondLst><p:childTnLst><p:par><p:cTn id="4" fill="hold"><p:stCondLst>'
            '<p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            + "".join(effects) +
            "</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>"
            "</p:childTnLst></p:cTn>"
            '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
            '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
            "</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>"
            + (f"<p:bldLst>{''.join(builds)}</p:bldLst>" if builds else "")
            + "</p:timing>"
        )
        el = etree.fromstring(xml)
        sld = slide._element
        after = sld.find(qn("p:transition"))
        if after is None:
            after = sld.find(f"{{{NS_MC}}}AlternateContent")
        if after is None:
            after = sld.find(qn("p:clrMapOvr"))
        after.addnext(el)

    # ------------------------------------------------------------ assembly

    def _metadata(self):
        c = self.d.get("cover") or {}
        cp = self.prs.core_properties
        cp.author = cp.last_modified_by = c.get("student", "")
        cp.title = c.get("title", "")
        cp.subject = c.get("subject", "")
        cp.comments = cp.keywords = cp.category = ""
        now = dt.datetime.now(dt.timezone.utc).replace(microsecond=0, tzinfo=None)
        cp.created = cp.modified = now
        cp.revision = 1
        pkg = self.prs.part.package
        for rid, rel in list(pkg._rels.items()):
            if rel.reltype == RT.THUMBNAIL:
                pkg._rels.pop(rid)

    def build(self) -> Presentation:
        slides = self.d.get("slides") or []
        sections = [s for s in slides if s.get("layout") == "section"]
        built = [self.title_slide()]
        if len(sections) >= 2:
            built.append(self.agenda_slide([s["title"] for s in sections]))
        sec_no = 0
        for sl in slides:
            lay = sl.get("layout")
            if lay == "section":
                sec_no += 1
                out = self.section_slide(sl, sec_no)
            elif lay == "two_column":
                out = self.two_column_slide(sl)
            elif lay == "stats":
                out = self.stats_slide(sl)
            elif lay == "table":
                out = self.table_slide(sl)
            elif lay == "quote":
                out = self.quote_slide(sl)
            elif lay == "timeline":
                out = self.timeline_slide(sl)
            else:
                out = self.bullets_slide(sl)
            if sl.get("notes"):
                out[0].notes_slide.notes_text_frame.text = sl["notes"]
            built.append(out)
        built.append(self.closing_slide())
        for slide, names in built:
            self.transition(slide)
            self.animate(slide, names)
        self._metadata()
        return self.prs


def build_pptx(deck: dict, out_path: str | Path) -> Path:
    out_path = Path(out_path)
    DeckBuilder(deck).build().save(str(out_path))
    return out_path

