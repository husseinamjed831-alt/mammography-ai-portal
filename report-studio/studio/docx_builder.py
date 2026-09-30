"""Build a polished academic report (.docx) from a structured dict.

The input shape is documented in README.md ("صيغة JSON"). The builder is
bilingual: Arabic reports are laid out right-to-left end to end (sections,
paragraphs, tables, runs), English reports left-to-right.

WordprocessingML is strict about child-element order, and Word refuses to
open a file that gets it wrong even when LibreOffice is forgiving, so every
low-level property goes through `_put()` with the schema order spelled out.
"""

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt

from .fonts import embed_fonts
from .themes import Theme, get_theme

BODY_FONT = "Amiri"
HEAD_FONT = "Tajawal"
DISPLAY_FONT = "Tajawal ExtraBold"

# ---------------------------------------------------------------- xml order

PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
    "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
    "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
    "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
    "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
    "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
    "sectPr", "pPrChange",
]
RPR_ORDER = [
    "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike",
    "dstrike", "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid",
    "vanish", "webHidden", "color", "spacing", "w", "kern", "position", "sz", "szCs",
    "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs",
    "em", "lang", "eastAsianLayout", "specVanish", "oMath",
]
TBLPR_ORDER = [
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
    "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription",
]
TCPR_ORDER = [
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
    "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
]
TRPR_ORDER = [
    "cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit",
    "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden",
]
SECTPR_ORDER = [
    "headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz",
    "pgMar", "paperSrc", "pgBorders", "lnNumType", "pgNumType", "cols", "formProt",
    "vAlign", "noEndnote", "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid",
    "printerSettings", "sectPrChange",
]
BORDER_ORDER = ["top", "left", "start", "bottom", "right", "end", "insideH", "insideV"]


def _el(tag: str, **attrs) -> OxmlElement:
    el = OxmlElement(f"w:{tag}")
    for k, v in attrs.items():
        el.set(qn(f"w:{k}"), str(v))
    return el


def _put(parent, child, order: list[str]):
    """Insert `child` into `parent` at its schema position, replacing any twin."""
    local = child.tag.split("}")[1]
    for existing in parent.findall(child.tag):
        parent.remove(existing)
    idx = order.index(local)
    later = {qn(f"w:{t}") for t in order[idx + 1 :]}
    for i, sib in enumerate(parent):
        if sib.tag in later:
            parent.insert(i, child)
            return child
    parent.append(child)
    return child


def _borders(container_tag: str, spec: dict[str, tuple[str, int, str] | None]):
    """spec: side -> (style, eighths-of-a-point, colour) or None for 'nil'."""
    box = _el(container_tag)
    for side in BORDER_ORDER:
        if side not in spec:
            continue
        v = spec[side]
        if v is None:
            box.append(_el(side, val="nil"))
        else:
            style, size, color = v
            box.append(_el(side, val=style, sz=size, space=0, color=color))
    return box


# ---------------------------------------------------------------- text utils

ARABIC_RE = re.compile(r"[؀-ۿݐ-ݿﭐ-﷿ﹰ-﻿]")
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")

ORDINALS_AR = [
    "أولاً", "ثانياً", "ثالثاً", "رابعاً", "خامساً", "سادساً", "سابعاً", "ثامناً",
    "تاسعاً", "عاشراً", "حادي عشر", "ثاني عشر", "ثالث عشر", "رابع عشر", "خامس عشر",
]

LABELS = {
    "ar": {
        "toc": "المحتويات", "abstract": "الملخص", "keywords": "الكلمات المفتاحية",
        "references": "المصادر والمراجع", "report_on": "تقرير بعنوان",
        "student": "إعداد الطالب", "subject": "المادة", "stage": "المرحلة",
        "supervisor": "بإشراف", "year": "العام الدراسي", "table": "جدول",
        "page": "الصفحة", "republic": "جمهورية العراق",
        "ministry": "وزارة التعليم العالي والبحث العلمي", "intro_label": "المقدمة",
    },
    "en": {
        "toc": "Contents", "abstract": "Abstract", "keywords": "Keywords",
        "references": "References", "report_on": "A report on",
        "student": "Prepared by", "subject": "Course", "stage": "Year / Stage",
        "supervisor": "Supervised by", "year": "Academic year", "table": "Table",
        "page": "Page", "republic": "Republic of Iraq",
        "ministry": "Ministry of Higher Education and Scientific Research",
        "intro_label": "Introduction",
    },
}


def has_arabic(text: str) -> bool:
    return bool(ARABIC_RE.search(text or ""))


# ---------------------------------------------------------------- builder


class ReportBuilder:
    def __init__(self, report: dict, toc_pages: list[int] | None = None):
        self.r = report
        self.lang = "en" if report.get("lang") == "en" else "ar"
        self.rtl = self.lang == "ar"
        self.L = LABELS[self.lang]
        self.theme: Theme = get_theme(report.get("theme"))
        self.opts = {
            "toc": True,
            "numbered_headings": True,
            "ministry_header": True,
            "page_border": self.theme.page_border,
            **(report.get("options") or {}),
        }
        self.toc_pages = toc_pages
        self.headings: list[tuple[int, str]] = []  # (level, display text) in order
        self.table_no = 0
        self.doc = Document()

    # ------------------------------------------------------------ primitives

    def _rpr(self, run, *, size=None, bold=False, color=None, font=None, italic=False, rtl=None):
        rPr = run._r.get_or_add_rPr()
        font = font or BODY_FONT
        _put(rPr, _el("rFonts", ascii=font, hAnsi=font, cs=font, eastAsia=font), RPR_ORDER)
        if bold:
            _put(rPr, _el("b"), RPR_ORDER)
            _put(rPr, _el("bCs"), RPR_ORDER)
        if italic:
            _put(rPr, _el("i"), RPR_ORDER)
            _put(rPr, _el("iCs"), RPR_ORDER)
        if color:
            _put(rPr, _el("color", val=color), RPR_ORDER)
        if size:
            hp = int(round(size * 2))
            _put(rPr, _el("sz", val=hp), RPR_ORDER)
            _put(rPr, _el("szCs", val=hp), RPR_ORDER)
        if self.rtl if rtl is None else rtl:
            _put(rPr, _el("rtl"), RPR_ORDER)

    def _ppr(self, p, *, align=None, before=0, after=6, line=None, keep_next=False,
             rtl=None, ind=None, border_bottom=None, shade=None, page_break=False,
             widow=True, outline=None):
        pPr = p._p.get_or_add_pPr()
        rtl = self.rtl if rtl is None else rtl
        if keep_next:
            _put(pPr, _el("keepNext"), PPR_ORDER)
            _put(pPr, _el("keepLines"), PPR_ORDER)
        if page_break:
            _put(pPr, _el("pageBreakBefore"), PPR_ORDER)
        if widow:
            _put(pPr, _el("widowControl"), PPR_ORDER)
        if border_bottom:
            style, size, color, space = border_bottom
            bdr = _el("pBdr")
            bdr.append(_el("bottom", val=style, sz=size, space=space, color=color))
            _put(pPr, bdr, PPR_ORDER)
        if shade:
            _put(pPr, _el("shd", val="clear", color="auto", fill=shade), PPR_ORDER)
        if rtl:
            _put(pPr, _el("bidi"), PPR_ORDER)
        spacing = {"before": int(before * 20), "after": int(after * 20)}
        if line:
            spacing.update(line=int(line * 240), lineRule="auto")
        _put(pPr, _el("spacing", **spacing), PPR_ORDER)
        if ind:
            _put(pPr, _el("ind", **{k: int(v) for k, v in ind.items()}), PPR_ORDER)
        # Leading-edge alignment is the default for both directions, so it is
        # expressed by *omitting* <w:jc>; this sidesteps the left/right swap
        # Word applies inside bidi paragraphs.
        if align in ("center", "both"):
            _put(pPr, _el("jc", val=align), PPR_ORDER)
        elif align == "end":
            # Inside a bidi paragraph Word reads left/right logically.
            _put(pPr, _el("jc", val="right" if not rtl else "left"), PPR_ORDER)
        if outline is not None:
            _put(pPr, _el("outlineLvl", val=outline), PPR_ORDER)

    def para(self, container, text="", *, size=14, bold=False, color=None, font=None,
             italic=False, rtl=None, style=None, rich=True, **ppr):
        p = container.add_paragraph(style=style) if style else container.add_paragraph()
        self._ppr(p, rtl=rtl, **ppr)
        if text:
            self.runs(p, text, size=size, bold=bold, color=color or self.theme.ink,
                      font=font, italic=italic, rtl=rtl, rich=rich)
        return p

    def runs(self, p, text, *, size, bold=False, color=None, font=None, italic=False,
             rtl=None, rich=True):
        parts = BOLD_RE.split(text) if rich else [text]
        for i, chunk in enumerate(parts):
            if not chunk:
                continue
            strong = bold or (i % 2 == 1)
            run = p.add_run(chunk)
            self._rpr(run, size=size, bold=strong, color=color, font=font, italic=italic, rtl=rtl)
        return p

    def field(self, p, instr: str, *, size, color, bold=False):
        def r(el):
            run = p.add_run()
            self._rpr(run, size=size, color=color, bold=bold)
            run._r.append(el)

        r(_el("fldChar", fldCharType="begin"))
        instr_el = OxmlElement("w:instrText")
        instr_el.set(qn("xml:space"), "preserve")
        instr_el.text = f" {instr} "
        r(instr_el)
        r(_el("fldChar", fldCharType="separate"))
        run = p.add_run("1")
        self._rpr(run, size=size, color=color, bold=bold)
        r(_el("fldChar", fldCharType="end"))

    def page_break(self):
        p = self.doc.add_paragraph()
        self._ppr(p, after=0)
        p.add_run().add_break(WD_BREAK.PAGE)

    # tables ----------------------------------------------------------------

    def table(self, container, rows: int, cols: int, *, widths_cm: list[float],
              borders: dict | None = None, center=False):
        t = container.add_table(rows=rows, cols=cols)
        tblPr = t._tbl.tblPr
        if self.rtl:
            _put(tblPr, _el("bidiVisual"), TBLPR_ORDER)
        total = sum(widths_cm)
        _put(tblPr, _el("tblW", w=int(Cm(total).twips), type="dxa"), TBLPR_ORDER)
        if center:
            _put(tblPr, _el("jc", val="center"), TBLPR_ORDER)
        _put(tblPr, _borders("tblBorders", borders or {
            s: None for s in ("top", "left", "bottom", "right", "insideH", "insideV")
        }), TBLPR_ORDER)
        _put(tblPr, _el("tblLayout", type="fixed"), TBLPR_ORDER)
        mar = _el("tblCellMar")
        for side, v in (("top", 60), ("left", 110), ("bottom", 60), ("right", 110)):
            mar.append(_el(side, w=v, type="dxa"))
        _put(tblPr, mar, TBLPR_ORDER)
        for i, w in enumerate(widths_cm):
            t.columns[i].width = Cm(w)
            for cell in t.columns[i].cells:
                cell.width = Cm(w)
        # python-docx seeds each cell with an empty paragraph; format it.
        for row in t.rows:
            for cell in row.cells:
                self._ppr(cell.paragraphs[0], after=0)
        return t

    def cell_style(self, cell, *, fill=None, borders=None, valign="center", margins=None):
        tcPr = cell._tc.get_or_add_tcPr()
        if borders:
            _put(tcPr, _borders("tcBorders", borders), TCPR_ORDER)
        if fill:
            _put(tcPr, _el("shd", val="clear", color="auto", fill=fill), TCPR_ORDER)
        if margins:
            mar = _el("tcMar")
            for side in ("top", "left", "bottom", "right"):
                if side in margins:
                    mar.append(_el(side, w=margins[side], type="dxa"))
            _put(tcPr, mar, TCPR_ORDER)
        _put(tcPr, _el("vAlign", val=valign), TCPR_ORDER)

    def cell_text(self, cell, text, **kw):
        p = cell.paragraphs[0]
        kw.setdefault("after", 0)
        ppr_keys = {"align", "before", "after", "line", "keep_next", "rtl", "ind"}
        self._ppr(p, **{k: v for k, v in kw.items() if k in ppr_keys})
        self.runs(p, text, **{k: v for k, v in kw.items() if k not in ppr_keys - {"rtl"}})
        return p

    # ------------------------------------------------------------ document

    def setup(self):
        d = self.doc
        styles = d.styles
        normal = styles["Normal"]
        rPr = normal.element.get_or_add_rPr()
        _put(rPr, _el("rFonts", ascii=BODY_FONT, hAnsi=BODY_FONT, cs=BODY_FONT, eastAsia=BODY_FONT), RPR_ORDER)
        _put(rPr, _el("sz", val=28), RPR_ORDER)
        _put(rPr, _el("szCs", val=28), RPR_ORDER)
        _put(rPr, _el("lang", val="en-US" if not self.rtl else "ar-IQ", bidi="ar-IQ"), RPR_ORDER)

        sec = d.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin = sec.bottom_margin = Cm(2.4)
        sec.left_margin = sec.right_margin = Cm(2.4)
        sec.header_distance = sec.footer_distance = Cm(1.1)
        sec.different_first_page_header_footer = True
        sectPr = sec._sectPr
        if self.rtl:
            _put(sectPr, _el("bidi"), SECTPR_ORDER)
        if self.opts["page_border"]:
            pb = _el("pgBorders", offsetFrom="page")
            for side in ("top", "left", "bottom", "right"):
                pb.append(_el(side, val="thinThickSmallGap", sz=24, space=24, color=self.theme.primary))
            _put(sectPr, pb, SECTPR_ORDER)

        self._header_footer(sec)
        self._metadata()

    def _metadata(self):
        c = self.r.get("cover") or {}
        cp = self.doc.core_properties
        who = c.get("student") or ""
        cp.author = who
        cp.last_modified_by = who
        cp.title = c.get("title") or ""
        cp.subject = c.get("subject") or ""
        cp.keywords = "، ".join(self.r.get("keywords") or [])
        cp.comments = ""
        cp.category = ""
        now = dt.datetime.now(dt.timezone.utc).replace(microsecond=0, tzinfo=None)
        cp.created = cp.modified = now
        cp.revision = 1
        # Drop the template's preview thumbnail.
        pkg = self.doc.part.package
        for rid, rel in list(pkg.rels.items()):
            if rel.reltype == RT.THUMBNAIL:
                pkg.rels.pop(rid)
                pkg.rels._target_parts_by_rId.pop(rid, None)

    def _header_footer(self, sec):
        c = self.r.get("cover") or {}
        th = self.theme
        head = sec.header.paragraphs[0]
        self._ppr(head, align="center", after=0,
                  border_bottom=("single", 6, th.accent, 4))
        label = " | ".join(x for x in (c.get("title"), c.get("subject")) if x)
        self.runs(head, label, size=9.5, color=th.muted, font=HEAD_FONT, rich=False)

        foot = sec.footer.paragraphs[0]
        self._ppr(foot, align="center", after=0)
        self.runs(foot, "— ", size=11, color=th.accent, font=HEAD_FONT, rich=False)
        self.field(foot, "PAGE", size=11, color=th.primary, bold=True)
        self.runs(foot, " —", size=11, color=th.accent, font=HEAD_FONT, rich=False)

    # cover -----------------------------------------------------------------

    def cover(self):
        c = self.r.get("cover") or {}
        th, L, d = self.theme, self.L, self.doc
        width = 16.2

        top_lines = []
        if self.opts["ministry_header"]:
            top_lines += [L["republic"], L["ministry"]]
        uni_line = " — ".join(x for x in (c.get("university"), c.get("college")) if x)
        if uni_line:
            top_lines.append(uni_line)
        if c.get("department"):
            top_lines.append(c["department"])

        if th.cover_style == "band":
            t = self.table(d, 1, 1, widths_cm=[width], center=True)
            cell = t.cell(0, 0)
            self.cell_style(cell, fill=th.primary, margins={"top": 320, "bottom": 320})
            first = True
            for i, line in enumerate(top_lines):
                p = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                big = i == len(top_lines) - (2 if c.get("department") else 1)
                self._ppr(p, align="center", after=2)
                self.runs(p, line, size=15 if big else 12.5, bold=True,
                          color="FFFFFF", font=HEAD_FONT, rich=False)
        else:
            for i, line in enumerate(top_lines):
                self.para(d, line, size=13.5, bold=True, color=th.primary, font=HEAD_FONT,
                          align="center", after=2, before=10 if i == 0 else 0, rich=False)
            self.para(d, "", align="center", after=0,
                      border_bottom=("double", 6, th.accent, 1))

        logo = c.get("logo_path")
        if logo and Path(logo).exists():
            p = self.para(d, "", align="center", before=18, after=0)
            p.add_run().add_picture(str(logo), height=Cm(3.2))
            gap = 30
        else:
            gap = 96

        self.para(d, L["report_on"], size=13, color=th.accent, font=HEAD_FONT,
                  align="center", before=gap, after=6, rich=False)
        self.para(d, c.get("title") or "", size=27 if self.rtl else 25, color=th.primary,
                  font=DISPLAY_FONT, align="center", after=6, line=1.15, rich=False)
        if c.get("subtitle"):
            self.para(d, c["subtitle"], size=14, color=th.muted, font=HEAD_FONT,
                      align="center", after=4, rich=False)
        self.para(d, "◆  ◆  ◆", size=10, color=th.accent, font=HEAD_FONT,
                  align="center", before=8, after=40, rich=False, rtl=False)

        info = [(L["student"], c.get("student")), (L["subject"], c.get("subject")),
                (L["stage"], c.get("stage")), (L["supervisor"], c.get("supervisor"))]
        info = [(k, v) for k, v in info if v]
        if info:
            rule = ("single", 4, "D0D4DA")
            t = self.table(d, len(info), 2, widths_cm=[4.0, 8.0], center=True,
                           borders={"top": None, "bottom": None, "left": None,
                                    "right": None, "insideV": None, "insideH": rule})
            for i, (k, v) in enumerate(info):
                self.cell_style(t.cell(i, 0), margins={"top": 90, "bottom": 90})
                self.cell_style(t.cell(i, 1), margins={"top": 90, "bottom": 90})
                self.cell_text(t.cell(i, 0), k, size=12.5, bold=True, color=th.accent, font=HEAD_FONT, rich=False)
                self.cell_text(t.cell(i, 1), v, size=14, bold=True, color=th.ink, rich=False)

        if c.get("academic_year"):
            self.para(d, f"{L['year']} {c['academic_year']}", size=12.5, color=th.muted,
                      font=HEAD_FONT, align="center", before=46, after=0, rich=False)
        self.page_break()

    # headings ----------------------------------------------------------------

    def h1(self, text: str, number: int | None = None, *, new_page=False):
        th = self.theme
        if number is not None and self.opts["numbered_headings"]:
            prefix = (ORDINALS_AR[number - 1] + ": ") if self.rtl and number <= len(ORDINALS_AR) else f"{number}. "
        else:
            prefix = ""
        shown = prefix + text
        p = self.para(self.doc, "", style="Heading 1", before=6 if new_page else 16, after=10,
                      keep_next=True, page_break=new_page, outline=0,
                      border_bottom=("single", 12, th.accent, 6))
        self.runs(p, shown, size=19 if self.rtl else 18, bold=True, color=th.primary,
                  font=HEAD_FONT, rich=False)
        self.headings.append((1, shown))

    def h2(self, text: str, number: str | None = None):
        shown = f"{number} {text}" if number and self.opts["numbered_headings"] else text
        p = self.para(self.doc, "", style="Heading 2", before=12, after=6, keep_next=True, outline=1)
        self.runs(p, shown, size=15 if self.rtl else 14, bold=True, color=self.theme.accent,
                  font=HEAD_FONT, rich=False)
        self.headings.append((2, shown))

    # blocks ----------------------------------------------------------------

    def body(self, text: str):
        # An English paragraph inside an Arabic report keeps LTR direction.
        rtl = self.rtl and (has_arabic(text) or not text.strip())
        self.para(self.doc, text, size=14 if self.rtl else 12, align="both",
                  after=8, line=1.2 if self.rtl else 1.15, rtl=rtl)

    def bullets(self, items: list[str], numbered=False):
        th = self.theme
        size = 14 if self.rtl else 12
        for i, item in enumerate(items, 1):
            p = self.para(self.doc, "", align="both", after=4, line=1.3,
                          ind={"start": 540, "hanging": 360})
            mark = f"{i}. " if numbered else "◂ " if self.rtl else "▸ "
            self.runs(p, mark, size=size - 1, bold=True, color=th.accent, font=HEAD_FONT, rich=False)
            self.runs(p, item, size=size, color=th.ink)
        # a little air after the list
        self.doc.paragraphs[-1].paragraph_format.space_after = Pt(10)

    def data_table(self, headers: list[str], rows: list[list[str]], caption: str | None):
        th = self.theme
        self.table_no += 1
        if caption:
            self.para(self.doc, f"{self.L['table']} ({self.table_no}): {caption}",
                      size=12 if self.rtl else 11, bold=True, color=th.primary, font=HEAD_FONT,
                      align="center", before=6, after=4, keep_next=True, rich=False)
        ncols = max(len(headers), max((len(r) for r in rows), default=0))
        widths = [16.2 / ncols] * ncols
        line = ("single", 4, "C9CED6")
        t = self.table(self.doc, 1 + len(rows), ncols, widths_cm=widths, center=True,
                       borders={"top": ("single", 12, th.primary), "bottom": ("single", 12, th.primary),
                                "left": None, "right": None, "insideH": line, "insideV": line})
        size = 12 if self.rtl else 10.5
        trPr = t.rows[0]._tr.get_or_add_trPr()
        _put(trPr, _el("tblHeader"), TRPR_ORDER)
        for j in range(ncols):
            cell = t.cell(0, j)
            self.cell_style(cell, fill=th.primary)
            self.cell_text(cell, headers[j] if j < len(headers) else "", size=size, bold=True,
                           color="FFFFFF", font=HEAD_FONT, align="center", rich=False)
        for i, row in enumerate(rows, 1):
            _put(t.rows[i]._tr.get_or_add_trPr(), _el("cantSplit"), TRPR_ORDER)
            for j in range(ncols):
                cell = t.cell(i, j)
                self.cell_style(cell, fill=th.soft if i % 2 == 0 else None)
                txt = row[j] if j < len(row) else ""
                self.cell_text(cell, txt, size=size, color=th.ink, align="center",
                               rtl=self.rtl and (has_arabic(txt) or not txt))
        self.para(self.doc, "", after=6)

    def callout(self, title: str | None, text: str):
        th = self.theme
        t = self.table(self.doc, 1, 1, widths_cm=[16.2], center=True,
                       borders={"top": ("single", 24, th.accent), "left": None, "right": None,
                                "bottom": None, "insideH": None, "insideV": None})
        cell = t.cell(0, 0)
        self.cell_style(cell, fill=th.soft, margins={"top": 140, "bottom": 140, "left": 220, "right": 220})
        size = 13 if self.rtl else 11.5
        if title:
            self.cell_text(cell, title, size=size, bold=True, color=th.primary, font=HEAD_FONT,
                           after=3, rich=False)
            p = cell.add_paragraph()
        else:
            p = cell.paragraphs[0]
        self._ppr(p, align="both", after=0, line=1.3)
        self.runs(p, text, size=size, color=th.ink)
        self.para(self.doc, "", after=6)

    def quote(self, text: str, source: str | None = None):
        th = self.theme
        size = 14 if self.rtl else 12
        self.para(self.doc, "", align="center", before=4, after=0,
                  ind={"start": 2600, "end": 2600}, border_bottom=("single", 8, th.accent, 1))
        self.para(self.doc, text, size=size + 1, color=th.primary, italic=not self.rtl,
                  align="center", before=6, after=2, line=1.35, ind={"start": 700, "end": 700})
        if source:
            self.para(self.doc, f"— {source}", size=size - 2, color=th.muted, align="center",
                      after=10, rich=False)

    def blocks(self, blocks: list[dict]):
        for b in blocks or []:
            kind = b.get("type")
            if kind == "paragraph":
                self.body(b.get("text", ""))
            elif kind in ("bullets", "numbered"):
                self.bullets(b.get("items") or [], numbered=kind == "numbered")
            elif kind == "table":
                self.data_table(b.get("headers") or [], b.get("rows") or [], b.get("caption"))
            elif kind == "callout":
                self.callout(b.get("title"), b.get("text", ""))
            elif kind == "quote":
                self.quote(b.get("text", ""), b.get("source"))

    # toc -------------------------------------------------------------------

    def planned_toc(self) -> list[tuple[int, str]]:
        """Headings the body will produce, in order (used to size the TOC)."""
        out = []
        if self.r.get("abstract"):
            out.append((1, self.L["abstract"]))
        for i, s in enumerate(self.r.get("sections") or [], 1):
            out.append((1, s["heading"]))
            for sub in s.get("subsections") or []:
                out.append((2, sub["heading"]))
        if self.r.get("references"):
            out.append((1, self.L["references"]))
        return out

    def toc(self):
        th, d = self.theme, self.doc
        entries = self.planned_toc()
        show_h2 = len(entries) <= 24
        self.para(d, self.L["toc"], size=22, color=th.primary, font=DISPLAY_FONT,
                  align="center", before=4, after=18, rich=False)
        visible = [(i, lvl) for i, (lvl, _) in enumerate(entries) if lvl == 1 or show_h2]
        t = self.table(d, len(visible), 2, widths_cm=[14.4, 1.8], center=True,
                       borders={"top": None, "left": None, "right": None, "insideV": None,
                                "bottom": ("dotted", 6, "B8BEC7"), "insideH": ("dotted", 6, "B8BEC7")})
        # Titles are filled in body order later, so keep a handle.
        self._toc_table = t
        self._toc_visible = visible
        for row, (idx, lvl) in enumerate(visible):
            left, right = t.cell(row, 0), t.cell(row, 1)
            self.cell_style(left, margins={"top": 70, "bottom": 70})
            self.cell_style(right, margins={"top": 70, "bottom": 70})
            page = "—"
            if self.toc_pages and idx < len(self.toc_pages) and self.toc_pages[idx]:
                page = str(self.toc_pages[idx])
            title = entries[idx][1]  # replaced with numbered text in fill_toc()
            ind = {"start": 0} if lvl == 1 else {"start": 480}
            self.cell_text(left, title, size=13.5 if lvl == 1 else 12.5, bold=lvl == 1,
                           color=th.ink if lvl == 1 else th.muted,
                           font=HEAD_FONT, ind=ind, rich=False)
            self.cell_text(right, page, size=12.5, bold=True, color=th.primary, font=HEAD_FONT,
                           align="end", rich=False)
        self.page_break()

    def fill_toc(self):
        """Swap planned titles for the numbered text actually rendered."""
        if not hasattr(self, "_toc_table"):
            return
        for row, (idx, _) in enumerate(self._toc_visible):
            if idx >= len(self.headings):
                continue
            p = self._toc_table.cell(row, 0).paragraphs[0]
            runs = p.runs
            if runs:
                runs[0].text = self.headings[idx][1]
                for extra in runs[1:]:
                    extra._r.getparent().remove(extra._r)

    # body ------------------------------------------------------------------

    def build(self) -> Document:
        self.setup()
        self.cover()
        if self.opts["toc"]:
            self.toc()
        first = True
        if self.r.get("abstract"):
            self.h1(self.L["abstract"], new_page=not first)
            first = False
            self.body(self.r["abstract"])
            kws = self.r.get("keywords") or []
            if kws:
                sep = "، " if self.rtl else ", "
                self.para(self.doc, f"**{self.L['keywords']}:** " + sep.join(kws),
                          size=13 if self.rtl else 11.5, after=10, align="both")
        for i, s in enumerate(self.r.get("sections") or [], 1):
            self.h1(s["heading"], i, new_page=first is False and s.get("new_page", False))
            first = False
            self.blocks(s.get("blocks"))
            for j, sub in enumerate(s.get("subsections") or [], 1):
                self.h2(sub["heading"], f"{i}.{j}")
                self.blocks(sub.get("blocks"))
        refs = self.r.get("references") or []
        if refs:
            self.h1(self.L["references"], new_page=True)
            for k, ref in enumerate(refs, 1):
                rtl = self.rtl and has_arabic(ref)
                p = self.para(self.doc, "", align="both" if rtl else None, after=6, line=1.2,
                              rtl=rtl, ind={"start": 480, "hanging": 480})
                self.runs(p, f"[{k}] ", size=11.5, bold=True, color=self.theme.accent,
                          font=HEAD_FONT, rtl=rtl, rich=False)
                self.runs(p, ref, size=13 if rtl else 11, color=self.theme.ink, rtl=rtl, rich=False)
        self.fill_toc()
        embed_fonts(self.doc)
        return self.doc


def build_docx(report: dict, out_path: str | Path, toc_pages: list[int] | None = None) -> list[tuple[int, str]]:
    b = ReportBuilder(report, toc_pages)
    b.build().save(str(out_path))
    return b.headings
